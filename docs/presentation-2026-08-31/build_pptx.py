"""
build_pptx.py
Rebuilds the review deck *inside the college's own PowerPoint template*.

Why this and not the HTML builder. The panel's template is a real .pptx with its
own theme, header artwork, footer and slide-number placeholders. Re-creating
that look in HTML is guesswork; opening the file and rewriting its content is
not. So every slide here is cloned from a slide that already exists in the
template -- a bullet slide, a table slide, a picture slide -- which carries the
theme, the footer and the numbering with it, and only the text and pictures are
replaced.

Content is the same as PRESENTATION.md; that file stays the readable source of
truth and this one is the thing you present from.

    uv run python docs/presentation-2026-08-31/build_pptx.py

Output: deck.pptx in this folder. For the PDF, convert the pptx itself rather
than re-rendering the content some other way -- that is the only version that
is still the college's template:

    nix run nixpkgs#libreoffice-fresh -- --headless --convert-to pdf \
        --outdir docs/presentation-2026-08-31 docs/presentation-2026-08-31/deck.pptx
"""

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = Path.home() / ".claude/uploads/a7495d6b-d005-4214-b2ed-0fc89e3fc78f/ca8e7234-Document_from_Vaishakh"
OUT = HERE / "deck.pptx"

# The template's own geometry, read off the slides rather than guessed.
BODY_LEFT, BODY_TOP = Emu(609600), Emu(1066800)
BODY_W, BODY_H = Emu(10972800), Emu(5181600)
TABLE_STYLE = "{5940675A-B579-460E-94D1-54222C63F5DA}"

DARK = RGBColor(0x12, 0x30, 0x4F)

# Slides in the template used as donors for each kind of slide.
DONOR_BULLETS, DONOR_TABLE, DONOR_PICTURE = 2, 5, 9   # 0-based: slides 3, 6, 10


# --- content ---------------------------------------------------------------
# (kind, title, payload). Bullets are (text, level) or plain strings.

CONTENTS = [
    "Background of the Problem", "Problem Statement", "Motivation",
    "Literature Review", "Research Gaps", "Project Objectives",
    "Proposed Methodology", "System Architecture / Flow Diagram",
    "Implementation and Results", "Working System (screenshots)",
    "Technologies and Resources", "Project Schedule", "References",
]

BACKGROUND = [
    "Every server writes a log line for each login, process and network event; a "
    "small company produces millions a day and nobody reads them by hand.",
    "Security teams rely on a SIEM (Splunk, QRadar, Wazuh) that matches those lines "
    "against written rules.",
    "A rule matches text, not meaning: “Failed password for root” matches, the "
    "same event worded differently does not.",
    "Published surveys of log anomaly detection [R2, R7] evaluate on fixed datasets "
    "where the attacker is not trying to avoid the detector.",
]

EVIDENCE = [
    "Sigma rules: 0 of ~3,300 match an SSH authentication line "
    "(searched for Failed password, Accepted password, Invalid user, pam_unix).",
    ("No coverage at all for the commonest attack surface on Linux.", 1),
    "Wazuh caught 91% of a brute force’s lines — and 29% of the same attack "
    "reworded. Same behaviour, different words.",
    "Our detector: 97–100% of those lines in both conditions.",
    "Alert volume on one clean corpus: Wazuh 4,546 alerts, ours 1,415 incidents, "
    "for 9 real attacks — alert fatigue, measured on both sides.",
    ("All four measured by us on real OpenSSH logs; sources in the report.", 1),
]

PROBLEM = [
    "SIEM rules match the text of a log line, not its meaning, so an attacker who "
    "rewords an attack walks straight past them.",
    "And nobody reports how much detection is actually lost when the attacker is "
    "deliberately hiding, so the size of the problem is unknown.",
]

MOTIVATION = [
    "Analysts drown in alerts: a production SIEM raised 4,546 alerts for 9 real "
    "attacks on one clean corpus.",
    "Words change, behaviour does not — many failures from one place, or a login "
    "from two countries an hour apart, survive rewording.",
    "Semantic detection can follow that; string matching cannot.",
    "It has to run where the budget is: the whole system runs on an ordinary laptop, "
    "no GPU and no cloud account.",
    ("A GPU only speeds the written explanation up: 3.9 s versus 18.2 s per incident.", 1),
]

LIT_HEAD = ["Paper (journal, year)", "Method", "Finding", "Limitation"]
LIT_1 = [
    ["R1 — Almodovar et al., LogFiT\nIEEE Trans. Network & Service Mgmt, 2024",
     "BERT fine-tuned on log text, self-supervised",
     "Language models read logs without labelled data",
     "Clean benchmark logs only; no evading attacker"],
    ["R2 — Landauer et al., survey\nMachine Learning with Applications, 2023",
     "Survey of deep learning for log anomaly detection",
     "Maps every major method and dataset in the field",
     "Reports evaluation is almost always static and non-adversarial"],
    ["R3 — Ali et al.\nEmpirical Software Engineering, 2025",
     "Classical ML versus deep learning, measured",
     "Simpler models compete with deep models on logs",
     "Clean data only; no adapting attacker"],
    ["R4 — Wu, Li & Khomh\nEmpirical Software Engineering, 2023",
     "Compares ways of turning logs into vectors",
     "Representation choice changes results a lot; embeddings work well",
     "Does not test reworded or disguised attacks"],
]
LIT_2 = [
    ["R5 — Khan et al.\nEmpirical Software Engineering, 2024",
     "Measures how parsing quality affects detection",
     "Poor parsing quietly ruins detection accuracy",
     "Parsing studied on clean logs"],
    ["R6 — Pang et al., review\nACM Computing Surveys, 2021",
     "Field review of deep anomaly detection",
     "Standard definitions and method families",
     "General; not security-log specific"],
    ["R7 — Kheddar, survey\nInformation Fusion, 2025",
     "Transformers and LLMs for intrusion detection",
     "LLMs are moving into intrusion detection quickly",
     "Evaluation still on fixed public datasets"],
    ["R8 — Debicha et al., TAD\nFuture Generation Computer Systems, 2023",
     "Attackers actively evading a detector",
     "Evasion measurably breaks detectors",
     "Network traffic, not logs; no SIEM comparison"],
]

GAPS_HEAD = ["What the literature does", "Gap we work in"]
GAPS = [
    ["Rule-based SIEMs match fixed text patterns [R2, R7]",
     "They miss the same attack once it is reworded — we measured 91% → 29%"],
    ["Log anomaly models are evaluated on clean benchmark data [R1, R3, R4, R5]",
     "Nobody reports how much detection survives an attacker who is hiding"],
    ["Adversarial evasion is studied for network intrusion detection [R8]",
     "Not for log analysis, and never against a production SIEM side by side"],
]

OBJECTIVES = [
    "Parse raw Linux authentication and system logs into a structured, maskable form.",
    "Detect anomalies from the meaning of a line, using sentence embeddings and a "
    "per-user Isolation Forest, with no labelled training data.",
    "Group related anomalies into incidents and attach the matching MITRE ATT&CK "
    "technique from a lookup table, never from the language model.",
    "Explain each incident in plain English with a local LLM that runs offline.",
    "Measure this system and a production SIEM (Wazuh) on the same logs under eight "
    "camouflage strategies, and report the difference — including where we lose.",
    "Keep the whole system running on a laptop with no GPU and no cloud service.",
]

METHOD_LEAD = [
    "Read logs → parse and mask → embed the meaning → score each user → group into "
    "incidents → attach MITRE technique → explain in English → show the analyst.",
    "Camouflaged attacks are fed back in and both systems re-scored on every change.",
    "Isolation Forest over embeddings is standard; the measurement, not the detector, "
    "is what this project contributes.",
]

ARCH_LEAD = [
    "Three live sources: a tailed file, this machine’s journal, or a paced replay.",
    "The model is fitted once on history and only scores the stream after that.",
    "Streaming reproduces the batch result exactly: 751 lines → 130 events → 94 incidents.",
    "The Wazuh comparison track is a separate package the pipeline may never import.",
]

IMPL_LEAD = [
    "One writer appends every line to both systems in the same call, so neither can be "
    "handed different traffic.",
    "Our injector wrote the attack, so every line’s label is recorded as it goes out: "
    "ground truth is automatic.",
    "The same scoring code grades both systems.",
    "Built so far: foundation and 449 automated tests · benchmark platform with a real "
    "SIEM · live streaming, API and dashboard · the live bake-off.",
]

RESULTS_HEAD = ["Camouflage strategy", "Our drop", "Wazuh’s drop", "Difference"]
RESULTS = [
    ["Reworded (rephrase)", "−2.2", "−33.3", "+31.1 in our favour"],
    ["Mimicry / noise / volume / timing", "0.0", "0.0", "tie × 4"],
    ["IP rotation", "−11.1", "0.0", "−11.1 against us"],
    ["Slowed down", "−8.9", "0.0", "−8.9 against us"],
    ["Realistic combination", "−13.3", "0.0", "−13.3 against us"],
]
RESULTS_NOTES = [
    "1 win, 4 ties, 3 losses, pooled over 5 seeds on real OpenSSH logs. The registered "
    "threshold (≥25 points on at least one strategy, written into git before any "
    "result existed) is cleared by rewording.",
    "Live: rewording an impossible-travel attack took Wazuh from 3 of 3 instances to "
    "0 of 3, while we stayed at 3 of 3.",
    "Where we are weaker: precision 0.29 — 94 incidents for 9 real attacks; impossible "
    "travel is caught by a distance-over-time rule, not by the embeddings.",
]

SHOTS_1 = [
    ("dashboard-incidents.png",
     "Live dashboard — 751 lines seen, 751 parsed, 94 incidents, 19 notable"),
    ("incident-narrative.png",
     "One incident opened — events, MITRE techniques, and the local LLM’s narrative"),
]
SHOTS_2 = [
    ("bakeoff-grid.png",
     "The bake-off grid — the examiner picks the attack and the camouflage level"),
    ("bakeoff-scored.png",
     "Both systems scored live on the same firing — 3 of 3 instances each, 100% of lines"),
]

TECH = [
    "Software: Python 3.11 with uv · sentence-transformers (all-MiniLM-L6-v2) · "
    "scikit-learn Isolation Forest · Ollama + phi3 for the local narrative.",
    "Comparison: the real Wazuh 4.9 manager in Docker, running its own rules.",
    "Interfaces and testing: FastAPI service · Streamlit dashboard · 449 pytest tests · "
    "GitHub Actions on every push.",
    "Data: LogHub OpenSSH (655,147 real lines, contains real internet attacks), LogHub "
    "Linux (99.4% parsed), and a 751-line synthetic corpus with full ground truth.",
    "Hardware, minimum and fully supported: any laptop, no GPU, 8 GB RAM.",
    ("Optional NVIDIA GPU only speeds the written explanation up (3.9 s vs 18.2 s).", 1),
    ("Demonstration setup: two laptops — one attacking and running the SIEM, one detecting.", 1),
]

GANTT_HEAD = ["Task", "Month 1", "Month 2", "Month 3", "Month 4"]
GANTT = [
    ["Literature review", "✓", "", "", ""],
    ["Requirement analysis", "✓", "✓", "", ""],
    ["Design", "", "✓", "", ""],
    ["Development", "", "✓", "", ""],
    ["Testing", "", "✓", "", ""],
    ["Documentation", "", "✓", "", ""],
]
GANTT_NOTES = [
    "Done so far: foundation and test suite · benchmark platform with a real SIEM · the "
    "pre-registered robustness measurement · live streaming, API and dashboard · the "
    "live side-by-side bake-off.",
    "Next: sequence modelling to fix precision, then an entity graph for correlation, "
    "then attacks generated by real tools.",
]

REFERENCES = [
    "[1] C. Almodovar, F. Sabrina, S. Karimi and S. Azad, “LogFiT: Log anomaly detection "
    "using fine-tuned language models,” IEEE Trans. Network and Service Management, "
    "21(2), pp. 1715–1723, 2024. doi:10.1109/TNSM.2024.3358730",
    "[2] M. Landauer, S. Onder, F. Skopik and M. Wurzenberger, “Deep learning for anomaly "
    "detection in log data: A survey,” Machine Learning with Applications, 12, 100470, "
    "2023. doi:10.1016/j.mlwa.2023.100470",
    "[3] S. Ali, C. Boufaied, D. Bianculli, P. Branco and L. Briand, “A comprehensive study "
    "of machine learning techniques for log-based anomaly detection,” Empirical Software "
    "Engineering, 30(5), 129, 2025. doi:10.1007/s10664-025-10669-3",
    "[4] X. Wu, H. Li and F. Khomh, “On the effectiveness of log representation for "
    "log-based anomaly detection,” Empirical Software Engineering, 28(6), 137, 2023. "
    "doi:10.1007/s10664-023-10364-1",
    "[5] Z. Khan, D. Shin, D. Bianculli and L. Briand, “Impact of log parsing on deep "
    "learning-based anomaly detection,” Empirical Software Engineering, 29(6), 139, 2024. "
    "doi:10.1007/s10664-024-10533-w",
    "[6] G. Pang, C. Shen, L. Cao and A. van den Hengel, “Deep learning for anomaly "
    "detection: A review,” ACM Computing Surveys, 54(2), pp. 1–38, 2021. doi:10.1145/3439950",
    "[7] H. Kheddar, “Transformers and large language models for efficient intrusion "
    "detection systems: A comprehensive survey,” Information Fusion, 124, 103347, 2025. "
    "doi:10.1016/j.inffus.2025.103347",
    "[8] I. Debicha, R. Bauwens, T. Debatty, J.-M. Dricot, T. Kenaza and W. Mees, “TAD: "
    "Transfer learning-based multi-adversarial detection of evasion attacks against network "
    "intrusion detection systems,” Future Generation Computer Systems, 138, pp. 185–197, "
    "2023. doi:10.1016/j.future.2022.08.011",
]


# --- building blocks -------------------------------------------------------

def clone(prs: Presentation, donor_index: int):
    """Copy a template slide, shapes and all, and append it.

    Cloning rather than `add_slide` keeps the theme, the footer and the
    slide-number field exactly as the template has them; a fresh slide from the
    layout drops the artwork this template puts on every page.
    """
    donor = prs.slides[donor_index]
    new = prs.slides.add_slide(donor.slide_layout)
    for shape in list(new.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in donor.shapes:
        new.shapes._spTree.append(copy.deepcopy(shape._element))
    return new


def set_title(slide, text: str) -> None:
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.idx == 0:
            shape.text_frame.text = text
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(30)
            return


def body_placeholder(slide):
    """The content placeholder, i.e. the one that is neither title nor chrome."""
    for shape in slide.shapes:
        if not shape.is_placeholder or not shape.has_text_frame:
            continue
        kind = shape.placeholder_format.type
        if shape.placeholder_format.idx != 0 and str(kind) not in (
                "SLIDE_NUMBER (13)", "FOOTER (15)", "DATE (16)"):
            return shape
    return None


def drop_non_chrome(slide) -> None:
    """Strip everything but the title, footer and slide number."""
    for shape in list(slide.shapes):
        keep = shape.is_placeholder and (
            shape.placeholder_format.idx == 0
            or str(shape.placeholder_format.type) in
            ("SLIDE_NUMBER (13)", "FOOTER (15)", "DATE (16)"))
        if not keep:
            shape._element.getparent().remove(shape._element)


def fill_bullets(slide, items, size: int = 18) -> None:
    box = body_placeholder(slide)
    if box is None:
        box = slide.shapes.add_textbox(BODY_LEFT, BODY_TOP, BODY_W, BODY_H)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        text, level = item if isinstance(item, tuple) else (item, 0)
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = text
        para.level = level
        para.space_after = Pt(8)
        for run in para.runs:
            run.font.size = Pt(size - 2 if level else size)


def add_table(slide, header, rows, top=None, height=None, size: int = 12) -> None:
    top = top or Emu(1600200)
    height = height or Emu(4200000)
    shape = slide.shapes.add_table(len(rows) + 1, len(header),
                                   BODY_LEFT, top, BODY_W, height)
    table = shape.table
    # Borrow the template's own table style rather than inventing one.
    style = shape._element.graphic.graphicData.tbl.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tblPr")
    if style is not None:
        for child in list(style):
            if child.tag.endswith("tableStyleId"):
                child.text = TABLE_STYLE
                break
        else:
            from lxml import etree
            node = etree.SubElement(
                style, "{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
            node.text = TABLE_STYLE
    for column, label in enumerate(header):
        cell = table.cell(0, column)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xE6, 0xEE, 0xF8)
        cell.text = label
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(size)
                run.font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(size)


def add_picture(slide, image: Path, top=None, max_h=None) -> None:
    """Place a picture centred in the body area, scaled to fit."""
    from PIL import Image

    top = top or Emu(1200000)
    max_h = max_h or Emu(4900000)
    with Image.open(image) as img:
        ratio = img.width / img.height
    width = BODY_W
    height = Emu(int(width / ratio))
    if height > max_h:
        height = max_h
        width = Emu(int(height * ratio))
    left = Emu(int((Emu(12192000) - width) / 2))
    slide.shapes.add_picture(str(image), left, top, width, height)


def add_caption(slide, text: str, top) -> None:
    box = slide.shapes.add_textbox(BODY_LEFT, top, BODY_W, Emu(320000))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for para in frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x4A, 0x5A, 0x6B)


def shot_slide(prs, title: str, filename: str, caption: str):
    """One capture, as large as the body area allows.

    A projected screenshot at half width is unreadable from the back of a room,
    which is the whole reason it is on a slide.
    """
    from PIL import Image

    slide = clone(prs, DONOR_PICTURE)
    drop_non_chrome(slide)
    set_title(slide, title)
    path = HERE / "screenshots" / filename
    with Image.open(path) as img:
        ratio = img.width / img.height
    height = Emu(4560000)
    width = Emu(int(height * ratio))
    if width > BODY_W:
        width, height = BODY_W, Emu(int(BODY_W / ratio))
    left = Emu(int((12192000 - width) / 2))
    slide.shapes.add_picture(str(path), left, Emu(1180000), width, height)
    add_caption(slide, caption, Emu(1180000) + height + Emu(80000))
    return slide


def shots_slide(prs, title: str, shots) -> None:
    """Two captured screenshots side by side, each with its caption."""
    slide = clone(prs, DONOR_PICTURE)
    drop_non_chrome(slide)
    set_title(slide, title)
    from PIL import Image

    gap = Emu(200000)
    width = Emu(int((BODY_W - gap) / 2))
    left = BODY_LEFT
    for filename, caption in shots:
        path = HERE / "screenshots" / filename
        if not path.exists():
            continue
        with Image.open(path) as img:
            ratio = img.width / img.height
        height = Emu(int(width / ratio))
        slide.shapes.add_picture(str(path), left, Emu(1350000), width, height)
        box = slide.shapes.add_textbox(left, Emu(1350000) + height + Emu(60000),
                                       width, Emu(600000))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = caption
        for para in frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0x4A, 0x5A, 0x6B)
        left = Emu(left + width + gap)


def bullets_slide(prs, title, items, size=18):
    slide = clone(prs, DONOR_BULLETS)
    set_title(slide, title)
    fill_bullets(slide, items, size=size)
    return slide


def table_slide(prs, title, header, rows, lead=None, size=12):
    slide = clone(prs, DONOR_TABLE)
    drop_non_chrome(slide)
    set_title(slide, title)
    top = Emu(1200000)
    if lead:
        box = slide.shapes.add_textbox(BODY_LEFT, Emu(1120000), BODY_W, Emu(400000))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = lead
        for run in frame.paragraphs[0].runs:
            run.font.size = Pt(13)
        top = Emu(1560000)
    height = Emu(min(4600000, 520000 * (len(rows) + 1)))
    add_table(slide, header, rows, top=top, height=height, size=size)
    return slide


def picture_slide(prs, title, image, lead=None):
    slide = clone(prs, DONOR_PICTURE)
    drop_non_chrome(slide)
    set_title(slide, title)
    top = Emu(1150000)
    if lead:
        box = slide.shapes.add_textbox(BODY_LEFT, Emu(1060000), BODY_W, Emu(900000))
        frame = box.text_frame
        frame.word_wrap = True
        for index, text in enumerate(lead):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.text = "•  " + text
            para.space_after = Pt(2)
            for run in para.runs:
                run.font.size = Pt(12)
        top = Emu(2050000)
    add_picture(slide, image, top=top, max_h=Emu(6858000) - top - Emu(560000))
    return slide


def reorder(prs, keep_order) -> None:
    """Rewrite the slide order and drop everything not in it."""
    id_list = prs.slides._sldIdLst
    entries = {id(slide): entry for slide, entry
               in zip(prs.slides, list(id_list))}
    for entry in list(id_list):
        id_list.remove(entry)
    for slide in keep_order:
        id_list.append(entries[id(slide)])


def main() -> None:
    prs = Presentation(str(TEMPLATE))
    original = list(prs.slides)
    title_slide, thanks, qa = original[0], original[17], original[18]

    # The title slide keeps its artwork; only the roll-number block and date move.
    for shape in title_slide.shapes:
        if shape.has_text_frame and "Roll No" in shape.text_frame.text:
            frame = shape.text_frame
            frame.clear()
            for index, line in enumerate([
                    "Harsh Bangera (Roll No. 3)",
                    "Manas Chaudhari (Roll No. 17)",
                    "Krish Mendon (Roll No. 62)",
                    "Anandan Nair (Roll No. 68)",
                    "",
                    "Prof. Darshan Mhapsekar",
                    "31 August 2026"]):
                para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                para.text = line
                para.alignment = PP_ALIGN.CENTER
                for run in para.runs:
                    run.font.size = Pt(18)

    deck = [title_slide]
    deck.append(bullets_slide(prs, "Contents", CONTENTS, size=18))
    deck.append(bullets_slide(prs, "Background of the Problem", BACKGROUND, size=17))
    deck.append(bullets_slide(prs, "Background — What We Measured", EVIDENCE, size=16))
    deck.append(bullets_slide(prs, "Problem Statement", PROBLEM, size=20))
    deck.append(bullets_slide(prs, "Motivation", MOTIVATION, size=17))
    deck.append(table_slide(prs, "Literature Review (1 of 2)", LIT_HEAD, LIT_1,
                            lead="Journal papers only — IEEE Transactions and Web of "
                                 "Science / Scopus indexed journals. Full PDFs are "
                                 "attached with this presentation.", size=11))
    deck.append(table_slide(prs, "Literature Review (2 of 2)", LIT_HEAD, LIT_2,
                            lead="All eight evaluate on data where the attacker is not "
                                 "trying to evade the detector being tested.", size=11))
    deck.append(table_slide(prs, "Research Gaps", GAPS_HEAD, GAPS,
                            lead="Our gap: how much worse does each system get when the "
                                 "attacker tries to hide — measured on the same logs, at "
                                 "the same moment, for a research detector and a "
                                 "production SIEM together.", size=13))
    deck.append(bullets_slide(prs, "Project Objectives", OBJECTIVES, size=16))
    deck.append(picture_slide(prs, "Proposed Methodology",
                              HERE / "diagrams/methodology.png", lead=METHOD_LEAD))
    deck.append(picture_slide(prs, "System Architecture / Flow Diagram",
                              HERE / "diagrams/architecture.png", lead=ARCH_LEAD))
    deck.append(picture_slide(prs, "Implementation — Keeping the Comparison Fair",
                              HERE / "diagrams/bakeoff.png", lead=IMPL_LEAD))
    deck.append(table_slide(prs, "Results — Robustness Under Camouflage",
                            RESULTS_HEAD, RESULTS,
                            lead="Drop in instance detection when the attack is "
                                 "camouflaged; bigger is worse. 5 seeds, real OpenSSH "
                                 "logs, both systems on the same corpus.", size=13))
    deck.append(bullets_slide(prs, "Results — How to Read Them", RESULTS_NOTES, size=16))
    deck.append(shot_slide(
        prs, "Working System (1 of 3) — Live Detection",
        "dashboard-incidents.png",
        "The dashboard reading the running service: 751 lines seen, 751 parsed, "
        "94 incidents, 19 notable, with the stream's own health above the list."))
    deck.append(shot_slide(
        prs, "Working System (2 of 3) — One Incident Explained",
        "incident-narrative.png",
        "The events that made the incident, the source addresses, the MITRE "
        "techniques from the lookup table, and the narrative written by phi3 on "
        "this laptop."))
    shots_slide(prs, "Working System (3 of 3) — The Live Bake-Off", SHOTS_2)
    deck.append(prs.slides[-1])
    deck.append(bullets_slide(prs, "Technologies and Resources", TECH, size=15))
    gantt = table_slide(prs, "Project Schedule & Current Progress", GANTT_HEAD, GANTT,
                        size=13)
    box = gantt.shapes.add_textbox(BODY_LEFT, Emu(5200000), BODY_W, Emu(1200000))
    frame = box.text_frame
    frame.word_wrap = True
    for index, text in enumerate(GANTT_NOTES):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = "•  " + text
        for run in para.runs:
            run.font.size = Pt(12)
    deck.append(gantt)
    deck.append(bullets_slide(prs, "References", REFERENCES, size=11))
    deck.extend([thanks, qa])

    reorder(prs, deck)
    prs.save(str(OUT))
    print(f"wrote {OUT.relative_to(Path.cwd())} — {len(deck)} slides")


if __name__ == "__main__":
    main()
